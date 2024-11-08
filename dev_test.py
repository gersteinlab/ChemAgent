from tqdm import tqdm
import pandas as pd
from datetime import datetime
import numpy as np
import sys
from io import StringIO

import json
import argparse
import math
import run
import json
import pandas as pd

import re
from typing import List
from colorama import Fore
from tenacity import retry, stop_after_attempt

import argparse
from XAgent.agent.base_agent import BaseAgent
from XAgent.utils import RequiredAbilities
from XAgent.message_history import Message
from XAgent.memory_db_plan import PlanMemoryDB
from XAgent.logs import logger
from XAgent.data_structure.node import ToolNode
from XAgent.ai_functions import function_manager, objgenerator
from XAgent.config import CONFIG

from XAgent.memory_db import MemoryDB_clear
import sometest


EVA_PROMPT_SYS = """
You are a teacher grading a quiz.
"""

EVA_PROMPT = """
You are given a question, the correct answer of the question, and the student's answer. You are asked to score the student's answer as either CORRECT or INCORRECT, based on the standard answer.
Write out in a step by step manner your reasoning to be sure that your conclusion is correct. Avoid simply stating the correct answer at the outset.

Example Format:
**QUESTION**: question here
**CORRECT ANSWER**: correct standard answer of the question here
**STUDENT ANSWER**: student's answer here
**EXPLANATION**: step by step reasoning here
**GRADE**: CORRECT or INCORRECT here

Grade the student answers based ONLY on their factual accuracy. Ignore differences in punctuation and phrasing between the student answer and true answer. It is OK if the student answer contains more information than the true answer, as long as it does not contain any conflicting statements. Begin! 

**QUESTION**: {{query}}
**STANDARD SOLUTION**: {{fir_answer}}
**STUDENT PLAN**: {{sec_answer}}
**EXPLANATION**:
"""


EVA_PLAN_PROMPT = """
You are given a question, the standard solution of the question, and the student's plan to solve it. You are asked to score the student's plan as either CORRECT or INCORRECT, based on the standard solution.
Write out in a step by step manner your reasoning to be sure that your conclusion is correct. Avoid simply stating the correct answer at the outset.

Example Format:
**QUESTION**: question here
**STANDARD SOLUTION**: standard solution of the question here
**STUDENT PLAN**: student's plan to solve the question here
**EXPLANATION**: step by step reasoning here
**GRADE**: CORRECT or INCORRECT here

Grade the student plans based on whether they are similar to the standard solution and whether student plan can lead to the same conclusion as the standard solution. Begin! 

**QUESTION**: {{query}}
**STANDARD SOLUTION**: {{standard}}
**STUDENT PLAN**: {{plan}}
**EXPLANATION**:
"""


class EvaAgent(BaseAgent):

    abilities = set([])

    def parse(
        self,
        placeholders: dict = {},
        arguments: dict = None,
        functions=None,
        function_call=None,
        stop=None,
        additional_messages: List[Message] = [],
        *args,
        **kwargs,
    ):
        logger.typewriter_log(f"debug:  ", Fore.BLUE, "parsing")
        prompt_messages = self.fill_in_placeholders(placeholders)
        messages = prompt_messages

        # for i in messages:
        #     logger.typewriter_log(f"debug:  ", Fore.BLUE, str(i.raw()))

        if isinstance(messages[0], Message):
            messages = [message.raw() for message in messages]

        logger.typewriter_log(f"debug :  ", Fore.BLUE, "reaching for response.........")
        response = objgenerator.chatcompletion(
            messages=messages,
            functions=functions,
            function_call=function_call,
            stop=stop,
            *args,
            **kwargs,
        )

        # logger.typewriter_log(f"debug response GET!:  ", Fore.BLUE, response)

        return response["choices"][0]["message"]["content"], response["usage"]


def remove_not(x):
    match_number = re.compile("[\$]?\ *10\^[{]?\ *-?[0-9]+\ *[}]?\ *[\$]?")
    result = re.findall(match_number, x)
    if len(result) != 0:
        return re.split(match_number, x)[-1]
    return None


def evaluate_plan(query, standard, plan):
    prompt_messages = [
        Message(role="system", content=EVA_PROMPT_SYS),
        Message(role="user", content=EVA_PLAN_PROMPT),
    ]
    eva_success = False
    attempts = 0
    same_vote = 0
    diff_vote = 0
    is_it_same = 2  # 2 means evaluation wrong
    while not eva_success and attempts < 10:
        try:
            agent = EvaAgent(CONFIG, prompt_messages)
            evaluation, _ = agent.parse(
                placeholders={
                    "user": {
                        "query": query,
                        "standard": standard,
                        "plan": plan,
                    }
                },
                schema_validation=False,
            )
            assert "GRADE" in evaluation, "No 'GRADE' tag"
            assert "CORRECT" in evaluation, "hallucination, no answer given"
            whether_same = evaluation.split("CORRECT")[-2][-2:] + "CORRECT"
            logger.typewriter_log(
                f"Evaluation after 'GRADE' tag : ", Fore.YELLOW, whether_same
            )
            if whether_same[:2] != "IN":
                same_vote += 1
            if whether_same[:2] == "IN":
                diff_vote += 1
            if same_vote >= 3 or diff_vote >= 2:
                eva_success = True
        except AssertionError as e:
            print(f"AssertionError on attempt {attempts + 1}: {e}. Retrying...")
            attempts += 1
        except Exception as e:
            print(f"Attempt {attempts + 1}: An error occurred: {e}. Retrying...")
            attempts += 1
    if same_vote >= 3 and diff_vote <= 1:
        is_it_same = 1
    if diff_vote >= 2:
        is_it_same = 0
    return is_it_same


def evaluate_ans(query, first_answer, second_answer):
    prompt_messages = [
        Message(role="system", content=EVA_PROMPT_SYS),
        Message(role="user", content=EVA_PROMPT),
    ]
    eva_success = False
    attempts = 0
    same_vote = 0
    diff_vote = 0
    is_it_same = 2  # 2 means evaluation wrong
    while not eva_success and attempts < 10:
        try:
            agent = EvaAgent(CONFIG, prompt_messages)
            evaluation, _ = agent.parse(
                placeholders={
                    "user": {
                        "query": query,
                        "fir_answer": first_answer,
                        "sec_answer": second_answer,
                    }
                },
                schema_validation=False,
            )
            assert "GRADE" in evaluation, "No 'GRADE' tag"
            assert "CORRECT" in evaluation, "hallucination, no answer given"
            whether_same = evaluation.split("CORRECT")[-2][-2:] + "CORRECT"
            logger.typewriter_log(
                f"Evaluation after 'GRADE' tag : ", Fore.YELLOW, whether_same
            )
            if whether_same[:2] != "IN":
                same_vote += 1
            if whether_same[:2] == "IN":
                diff_vote += 1
            if same_vote >= 3 or diff_vote >= 2:
                eva_success = True
        except AssertionError as e:
            print(f"AssertionError on attempt {attempts + 1}: {e}. Retrying...")
            attempts += 1
        except Exception as e:
            print(f"Attempt {attempts + 1}: An error occurred: {e}. Retrying...")
            attempts += 1
    if same_vote >= 3 and diff_vote <= 1:
        is_it_same = 1
    if diff_vote >= 2:
        is_it_same = 0
    return is_it_same


def evaluate_no_llm(query, fir, sec):
    pattern = r"[-+]?[0-9]*\.?[0-9]+"
    if not re.findall(pattern, sec):
        return 0
    fir_numbers = float(re.findall(pattern, fir)[0])
    sec_numberss = []
    for i in re.findall(pattern, sec):
        sec_numberss.append(float(i))
    is_it_same = 2
    for sec_numbers in sec_numberss:
        try:
            if sec_numbers >= 1:
                is_it_same = (
                    1 if math.isclose(fir_numbers, sec_numbers, rel_tol=0.05) else 0
                )
                if is_it_same == 0:
                    is_it_same = (
                        1 if math.isclose(fir_numbers, float(str(format(sec_numbers,'.2e')).split('e')[0]), rel_tol=0.05) else 0
                    )
            else:
                is_it_same = (
                    1 if math.isclose(fir_numbers, sec_numbers, rel_tol=0.1) else 0
                )
                if is_it_same == 0:
                    is_it_same = (
                        1 if math.isclose(fir_numbers, float(str(format(sec_numbers,'.2e')).split('e')[0]), rel_tol=0.1) else 0
                    )
        except:
            is_it_same = 2
        if is_it_same == 1:
            break
    return is_it_same


def parse_args() -> argparse.Namespace:
    """
    Parse the command line arguments and return them as an argparse.Namespace object.

    Returns:
        argparse.Namespace: An object containing command line arguments and their values.
    """
    parser = argparse.ArgumentParser()

    # parser.add_argument("--wholetask", type=str, required=True, help="The whole task description.")
    # parser.add_argument("--wholesolution", type=str, default = "Not provided.", help="The whole task solution.")
    parser.add_argument(
        "--list_source",
        nargs="+",
        default=[
            "chemmc",
            "class",
            "diff",
            "fund",
            "matter",
            "quan",
            "stat",
            "thermo",
            "atkins",
            "calculus",
        ],
    )
    parser.add_argument(
        "--mode",
        choices=["dev", "test", "dev_and_test", "debug"],
        default="dev",
        help="mode type",
    )
    parser.add_argument(
        "--tool",
        default="python",
        help="use python as tool",
    )
    parser.add_argument(
        "--global_mode",
        action="store_true",
        help="whether share global memory on different subject.",
    )
    parser.add_argument(
        "--empty",
        action="store_true",
        help="whether clear memory and start from an empty one.",
    )
    parser.add_argument("--iter", type=int, default=1, help="run a test `iter` times")
    parser.add_argument(
        "--tag", default="none", help="other tag that will be saved in xlsx"
    )
    return parser.parse_args()


def dev(tool, subject):

    failsplit = 0
    tot = 0
    eva_failed = 0
    cor = 0
    with open("dataset/{}_sol.json".format(subject), encoding="utf-8") as json_file:
        problems = json.load(json_file)
        for problem_data in tqdm(problems):
            corr_answer = problem_data["answer_number"]
            solution = problem_data["solution"]
            unit_prob = problem_data["unit"]
            if remove_not(unit_prob):
                unit_prob = remove_not(unit_prob)
            whole_task = problem_data["problem_text"]
            if unit_prob != " " and unit_prob != None:
                whole_task = (
                    whole_task + "The unit of the answer should be " + unit_prob
                )
            whole_solution = (
                solution + " So my final answer is : " + corr_answer + " " + unit_prob
            )
            tasks_and_solutions = sometest.env(
                whole_task=whole_task, whole_solution=whole_solution
            )

            if "FAILED" in tasks_and_solutions[0]:
                logger.typewriter_log(
                    f"SplitSort Agent Error:  ", Fore.RED, tasks_and_solutions[0]
                )
                failsplit += 1
                continue
            for task_and_solution in tasks_and_solutions:
                logger.typewriter_log(
                    f"Splitted and sorted tasks:  --------------------", Fore.RED, " "
                )
                print(task_and_solution)
            for task_and_solution in tasks_and_solutions:
                tot += 1
                this_task = task_and_solution.split("SOLUTION:")[0] + "\n"
                this_solution = (
                    "Here is a correct solution trajectory that you can refer to. "
                    + task_and_solution.split("SOLUTION:")[1]
                )
                this_final_ans = "None"
                this_subtasks = "I do not know how to solve this task."
                if task_and_solution != tasks_and_solutions[-1]:
                    try:
                        this_final_ans, this_subtasks, this_knowledge, _ = run.run(
                            [
                                "--task",
                                this_task + this_solution,
                                "--insert",
                                "True",
                                "--tool",
                                tool,
                            ]
                        )
                    except:
                        this_final_ans = "error occurs when solving the task"
                        this_subtasks = "I do not know how to solve this task."

                    same_or_not = evaluate_plan(
                        this_task,
                        this_solution,
                        "My plan to solve this question is listed below: \n"
                        + this_subtasks,
                    )
                    if same_or_not == 2:
                        eva_failed += 1
                    if same_or_not == 1:
                        cor += 1

                        DB = PlanMemoryDB(init_mode=False)
                        DB.insert_sentence(
                            this_task.split("QUESTION:")[1], this_knowledge, "SUCCESS"
                        )
                        DB.save_vector()

                        print(this_knowledge)
                        print("_____insert plan______\n")
                    logger.typewriter_log(
                        f"This task:  ------------------------", Fore.RED, " "
                    )
                    print(this_task, "\n", this_subtasks)
                    logger.typewriter_log(
                        f"This solution:  ------------------------", Fore.RED, " "
                    )
                    print(this_solution)
                    logger.typewriter_log(
                        f"This answer:  ------------------------", Fore.RED, " "
                    )
                    print(this_final_ans)
                    logger.typewriter_log(
                        f"This task End  ------------------------", Fore.RED, " "
                    )
                if task_and_solution == tasks_and_solutions[-1]:
                    try:
                        this_final_ans, this_subtasks, this_knowledge, _ = run.run(
                            [
                                "--task",
                                this_task + this_solution.split(corr_answer)[0],
                                "--insert",
                                "True",
                                "--tool",
                                tool,
                            ]
                        )
                    except:
                        this_final_ans = "error occurs when solving the task"
                    if not this_final_ans or this_final_ans.strip() == "":
                        this_final_ans = "None"
                    print("final answer is: \n", this_final_ans, "\n")
                    print("correct answer is: \n", corr_answer + " " + unit_prob, "\n")
                    same_or_not = evaluate_ans(
                        this_task, corr_answer + " " + unit_prob, this_final_ans
                    )
                    if same_or_not == 2:
                        eva_failed += 1
                    if same_or_not == 1:
                        cor += 1

                        DB = PlanMemoryDB(init_mode=False)
                        DB.insert_sentence(
                            this_task.split("QUESTION:")[1], this_knowledge, "SUCCESS"
                        )
                        DB.save_vector()

                        print(this_knowledge)
                        print("!!!!!!!!!!!!inserted!!!!!!!!!!!!!!!!!")
                    print(tot, cor, eva_failed, (cor / (tot - eva_failed)), (cor / tot))


def test(tool, subject, tag):
    cor = 0
    tot = 0
    eva_failed = 0
    final_answers = []
    corr_answers = []
    solve_func = []
    corr_final_answers = []
    corr_corr_answers = []
    corr_solve = []
    visual_answer_pairs = []
    with open("dataset/{}.json".format(subject), encoding="utf-8") as json_file:
        problems = json.load(json_file)
        for problem_data in tqdm(problems):
            tot += 1
            corr_answer = problem_data["answer_number"]
            unit_prob = problem_data["unit"]
            unit_prob_print = problem_data["unit"]
            if remove_not(unit_prob):
                unit_prob = remove_not(unit_prob)
            problem_text_0 = (
                problem_data["problem_text"]
                ## + " The answer should be represented in float. "
                ## + " And the unit of the answer is "
                ## + unit_prob_question
                ## + "."
                ## + " You should provide your final answer in this unit (if required)."
            )
            if unit_prob != " " and unit_prob != None:
                problem_text_0 = (
                    problem_text_0 + "The unit of the answer should be " + unit_prob
                )
            final_answer_0 = "None"
            final_answer = "None"
            msg = "None"
            try:
                final_answer_0, _, _, msg = run.run(
                    ["--task", problem_text_0, "--insert", "False", "--tool", tool]
                )
                ## if unit_prob != " ":
                ##     problem_text_2 = 'Given an answer "'+ final_answer_0 + '". \nconvert it to this unit: `' + unit_prob + '`.'
                ##     final_answer,_ = run.run(
                ##         ["--task", problem_text_2, "--insert", "False", "--tool", tool]
                ##     )
                ## else:
                final_answer = final_answer_0
            except Exception as e:
                print(e)
                final_answer = "error occurs when solving the task"
            print("final answer is: \n", final_answer, "\n")
            print("correct answer is: \n", corr_answer + " " + unit_prob_print, "\n")
            if not final_answer or final_answer.strip() == "":
                final_answer = "None"
            final_answers.append(final_answer)
            solve_func.append(msg)
            corr_answers.append(corr_answer + " " + unit_prob_print)
            # same_or_not = evaluate(final_answer, corr_answer + " " + unit_prob)
            same_or_not = evaluate_no_llm(problem_text_0, corr_answer, final_answer)
            if same_or_not == 2:
                eva_failed += 1
            if same_or_not == 1:
                cor += 1
                corr_final_answers.append(final_answer)
                corr_solve.append(msg)
                corr_corr_answers.append(corr_answer + " " + unit_prob_print)
            print(tot, cor, eva_failed, (cor / (tot - eva_failed)), (cor / tot))
    print("total task's number: ", tot)
    print("correct solved tasks: ", cor)
    print("evaluation failed tasks: ", eva_failed)
    print(
        "accuracy (not consider evaluation failed ones): ", (cor / (tot - eva_failed))
    )
    print("accuracy (full): ", (cor / tot))
    for i in range(len(final_answers)):
        print(final_answers[i], ";", corr_answers[i], "\n")
        visual_answer_pairs.append(final_answers[i] + ";")
        visual_answer_pairs.append(corr_answers[i] + " <>")
    print(corr_final_answers)
    print(corr_corr_answers)

    with open("./results_{}/{}.txt".format(tool, subject), "a+", encoding="utf-8") as f:
        f.write(
            f"Config: img {str(CONFIG.image)},refine {str(CONFIG.refine)},score {str(CONFIG.score)}, eva {CONFIG.eva_model}\n"
        )
        f.write("total task's number: " + str(tot) + "\n")
        f.write("correct solved tasks: " + str(cor) + "\n")
        f.write("evaluation failed tasks: " + str(eva_failed) + "\n")
        f.write(
            "accuracy (not consider evaluation failed ones): "
            + str(cor / (tot - eva_failed))
            + "\n"
        )
        f.write("accuracy (full): " + str((cor / tot)) + "\n")
        f.write("\n\n-----------------------------------\n\n")
        f.write(
            "\n".join(
                [
                    (
                        final_answers[i]
                        + "  ;  "
                        + corr_answers[i]
                        + " ; "
                        + solve_func[i]
                    )
                    for i in range(len(final_answers))
                ]
            )
        )
        f.write("\n\n-----------------------------------\n\n")
        f.write(
            "\n".join(
                [
                    (
                        corr_final_answers[i]
                        + "  ;  "
                        + corr_corr_answers[i]
                        + " ; "
                        + corr_solve[i]
                    )
                    for i in range(len(corr_final_answers))
                ]
            )
        )
        f.write("\n\n-----------------------------------\n\n")

    data = [
        tot,
        cor,
        eva_failed,
        cor / (tot - eva_failed),
        cor / tot,
        datetime.now(),
        str(final_answers),
        str(corr_answers),
        str(corr_final_answers),
        str(corr_corr_answers),
        tool,
        subject,
        CONFIG.save_dir,
        tag,
        CONFIG.plan_save_dir,
        str(visual_answer_pairs),
    ]

    import openpyxl

    try:

        workbook = openpyxl.load_workbook(
            "results_excel.xlsx",
        )

        worksheet = workbook["sheet_now"]

        worksheet.append(data)

        workbook.save(
            "results_excel.xlsx",
        )
        workbook.close()
        print("Excel write success！")
    except FileNotFoundError:
        print("Excel file not found, create a new one")


if __name__ == "__main__":
    args = parse_args()
    if args.global_mode and args.empty:
        run.run(
            [
                "--task",
                "Empty task.\n",
                "--insert",
                "False",
                "--tool",
                args.tool,
                "--init",
                "True",
            ]
        )
    for source in args.list_source:
        if args.mode == "dev":
            if not args.global_mode and args.empty:
                run.run(
                    [
                        "--task",
                        "Empty task.\n",
                        "--insert",
                        "False",
                        "--tool",
                        args.tool,
                        "--init",
                        "True",
                    ]
                )
            dev(args.tool, source)
        if args.mode == "test":
            for i in range(args.iter):
                test(args.tool, source, args.tag)
        if args.mode == "dev_and_test":
            if not args.global_mode and args.empty:
                run.run(
                    [
                        "--task",
                        "Empty task.\n",
                        "--insert",
                        "False",
                        "--tool",
                        args.tool,
                        "--init",
                        "True",
                    ]
                )
            dev(args.tool, source)
            test(args.tool, source)
        if args.mode == "debug":
            pass


"""
USAGE=>

dev_test.py --list_source chemmc --mode test --tool simple

"""
